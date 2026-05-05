import pyautogui
import pyperclip
import time
from openpyxl import load_workbook
import sys

# ===== ARQUIVO ==== 
ARQUIVO = sys.argv[1] if len(sys.argv) > 1 else "planilha.xlsx"

# ===== COLUNAS =====
COL_CODIGO = "B"
COL_NOME = "C"
COL_STATUS = "D"

LINHA_INICIAL = 2

# ===== POSIÇÕES (AJUSTAR SE PRECISAR) =====
CAMPO_PESQUISA = (882, 100)
AREA_NOME_LISTA = (266, 374)
AREA_COPIAR_NOME = (514, 297)
BOTAO_SAIR_USUARIO = (760, 518)

# ===== TEMPOS =====
TEMPO_BUSCA = 1.5
TEMPO_ENTRAR = 1.0
TEMPO_SAIR = 0.8

pyautogui.PAUSE = 0.3


# ===== FUNÇÕES =====

def clicar(pos):
    pyautogui.click(pos[0], pos[1])


def escrever(texto):
    pyperclip.copy(str(texto))
    pyautogui.hotkey("ctrl", "v")


def limpar():
    pyautogui.hotkey("ctrl", "a")
    pyautogui.press("backspace")


def clicar_no_nome():
    clicar(AREA_NOME_LISTA)
    time.sleep(TEMPO_ENTRAR)


def sair_do_usuario():
    clicar(BOTAO_SAIR_USUARIO)
    time.sleep(TEMPO_SAIR)


def copiar_nome():
    for _ in range(2):  # tenta 2 vezes
        pyautogui.click(
            AREA_COPIAR_NOME[0],
            AREA_COPIAR_NOME[1],
            clicks=4,
            interval=0.08
        )
        time.sleep(0.2)

        pyautogui.hotkey("ctrl", "c")
        time.sleep(0.3)

        nome = pyperclip.paste().strip()

        if nome:
            return nome

    return ""


# ===== MAIN =====

def rodar():
    wb = load_workbook(ARQUIVO)
    ws = wb.active

    print("Abra o sistema e deixe na tela de busca.")
    input("Pressione ENTER para começar")

    linha = LINHA_INICIAL

    while True:
        codigo = ws[f"{COL_CODIGO}{linha}"].value

        if not codigo:
            break

        print(f"Linha {linha} -> {codigo}")

        try:
            # digita código
            clicar(CAMPO_PESQUISA)
            limpar()
            escrever(codigo)

            time.sleep(0.2)
            pyautogui.press("enter")
            time.sleep(TEMPO_BUSCA)

            # entra no usuário
            clicar_no_nome()

            # copia nome
            nome = copiar_nome()

            if nome:
                ws[f"{COL_NOME}{linha}"] = nome
                ws[f"{COL_STATUS}{linha}"] = "OK"
            else:
                ws[f"{COL_STATUS}{linha}"] = "NAO COPIOU"

            # sair do usuário
            sair_do_usuario()

        except Exception as e:
            ws[f"{COL_STATUS}{linha}"] = f"ERRO: {e}"

        wb.save(ARQUIVO)
        linha += 1
        time.sleep(1)

    print("Finalizado!")


if __name__ == "__main__":
    rodar()